from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = BASE_DIR / "data" / "run_sql_export.xlsx"
DEFAULT_OUTPUT_DIR = BASE_DIR / "output"

REQUIRED_SHEETS = [
    "ReportDataDictionary",
    "Time",
]

REPORT_DATA_SHEETS = [
    "ReportData_1",
    "ReportData_2",
    "ReportData_3",
]

KEY_VALUABLE_NAMES = [
    "Zone Ideal Loads Supply Air Total Cooling Energy",
    "Zone Ideal Loads Supply Air Total Heating Energy",
    "Zone Ideal Loads Zone Total Cooling Energy",
    "Zone Ideal Loads Zone Total Heating Energy",
    "Enclosure Windows Total Transmitted Solar Radiation Energy",
    "Zone Lights Total Heating Energy",
    "Zone Electric Equipment Total Heating Energy",
    "Zone People Total Heating Energy",
    "Zone Mean Air Temperature",
    "Zone Operative Temperature",
    "Zone Air Relative Humidity",
    "Surface Window Heat Gain Energy",
    "Surface Window Heat Loss Energy",
    "Surface Inside Face Temperature",
    "Surface Outside Face Temperature",
]

COOLING_NAME = "Zone Ideal Loads Supply Air Total Cooling Energy"
HEATING_NAME = "Zone Ideal Loads Supply Air Total Heating Energy"
HEAT_CONTRIBUTION_VARIABLES = {
    "people": "Zone People Total Heating Energy",
    "lighting": "Zone Lights Total Heating Energy",
    "equipment": "Zone Electric Equipment Total Heating Energy",
    "solar_transmission": (
        "Enclosure Windows Total Transmitted Solar Radiation Energy"
    ),
}

def parse_arg() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inspect Energyplus SQL exported Excel workbook"
    )
    parser.add_argument(
        "--input",
        default=str(DEFAULT_INPUT),
        help="Input Excel file path",
    )
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Output folder path",
    )
    return(parser.parse_args())

def resolve_path(path_text:str) -> Path:  # 变绝对路径
    path = Path(path_text)

    if not path.is_absolute():
        path = BASE_DIR / path

    return(path)

def read_sheet_as_dataframe(workbook,sheet_name:str) -> pd.DataFrame: # 将sheet页变成dataframe
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found:{sheet_name}")
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))   # 后续不嵌套list会更好

    if not rows:
        return(pd.DataFrame())

    header = list(rows[0])
    data = [list(row) for row in rows[1:]]

    return pd.DataFrame(data,columns=header)

def validate_sheets(sheet_names:list[str]) ->list[str]: # 看是否有缺的必须sheet页
    required = REQUIRED_SHEETS + REPORT_DATA_SHEETS
    missing = [sheet for sheet in required if sheet not in sheet_names]
    return missing

def find_key_variables(dictionary_df:pd.Dataframe) -> pd.DataFrame: #筛选关键指标的数据，剔除其他
    required_columns = [
         "ReportDataDictionaryIndex",
        "KeyValue",
        "Name",
        "Units",
        "ReportingFrequency",
    ]
    missing_columns = [
        column for column in required_columns if column not in dictionary_df
    ]
    if missing_columns:
        raise ValueError(
            "Missing columns in ReportDataDictionary:"
            +",".join(missing_columns)
        )
    key_variables = dictionary_df[
        dictionary_df["Name"].astype(str).isin(KEY_VALUABLE_NAMES)
    ].copy()

    key_variables = key_variables[required_columns]
    key_variables = key_variables.sort_values(
        by=["Name","KeyValue"],
        ignore_index =True  #抛弃原来的索引
    )
    return key_variables

def get_key_variable_indices(key_variables:pd.DataFrame) ->set[int]: # 找关键数据的index
    indices = pd.to_numeric(
        key_variables["ReportDataDictionaryIndex"],
        errors="coerce"  #讲错误数据改为Nan
    )
    indices = indices.dropna().astype(int)
    return set(indices)

def extract_selected_report_data_from_sheet(   # 将关键数据单独抽出一个dataframe
        workbook,
        sheet_name:str,
        key_variable_indices:set[int]
) -> pd.DataFrame:
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True) # 获得sheet页中的所有行
    header = next(rows) #将第一行做为表头

    column_index = {
        column_name:index
        for index,column_name in enumerate(header)
    }    #制作“一、列名与index对应字典”

    required_columns = [
        "ReportDataIndex",
        "TimeIndex",
        "ReportDataDictionaryIndex",
        "Value",
    ]
    selected_rows =[]

    for row in rows:
        dictionary_index = row[column_index["ReportDataDictionaryIndex"]] # 根据“一、列名与index对应字典”找到数值标签这一整列

        if dictionary_index in key_variable_indices:
            selected_rows.append(
                {
                    column:row[column_index[column]]
                    for column in required_columns
                }
            )
    return(pd.DataFrame(selected_rows,columns=required_columns))

# ------ 给简单数据报告表加上更多的信息 ------

def attach_variable_info(
        report_data:pd.DataFrame,
        key_variables:pd.DataFrame,
) -> pd.DataFrame:
    variable_info = key_variables[
        [
            "ReportDataDictionaryIndex",
            "KeyValue",
            "Name",
            "Units",
        ]
    ].copy()

    merged_data = report_data.merge(
        variable_info,
        on="ReportDataDictionaryIndex", # 以哪一列为对照，这一列相同的俩个表的数据会放在同一行
        how="left", # 左边表格的行全部保留，右边表格只负责贴信息，贴不上的就空着。
    )

    return(merged_data)

# ------ 给数据报告上加上时间信息 ------
def attach_time_info(
        report_data:pd.DataFrame,
        time_table:pd.DataFrame
) -> pd.DataFrame:
    merged_data = report_data.merge(
        time_table,
        on="TimeIndex",
        how="left",
    )
    return merged_data

def find_missing_time_indices(time_df:pd.DataFrame) -> list[int]:
    if "TimeIndex" not in time_df.columns:
        return []

    time_indices = pd.to_numeric(time_df["TimeIndex"],errors="coerce")
    time_indices = time_indices.dropna().astype(int)

    if time_indices.empty:
        return []

    start = int(time_indices.min())
    end = int(time_indices.max())

    expected = set(range(start,end + 1))
    actual = set(time_indices)
    return sorted(expected - actual)

# ------ 将time时刻表转换为带有标准时间戳的Dataframe ------
def prepare_time_table(time_df:pd.DataFrame) ->pd.DataFrame:
    required_columns = [
        "TimeIndex",
        "Year",
        "Month",
        "Day",
        "Hour",
        "Minute",
    ]
    missing_columns = [
        column for column in required_columns if column not in time_df.columns
    ]
    if missing_columns:
        raise ValueError(
            "Missing columns in time sheet:"+",".join(missing_columns)
        )
    time_table = time_df[required_columns].copy()
    for column in required_columns:
        time_table[column] = pd.to_numeric(time_table[column],errors="coerce")

    date_parts = pd.DataFrame(
        {
            "year":time_table["Year"],
            "month":time_table["Month"],
            "day":time_table["Day"],
        }
    )
    base_date = pd.to_datetime(date_parts,errors="coerce")
    time_table["datetime"] = (
        base_date
        + pd.to_timedelta(time_table["Hour"],unit="h")
        + pd.to_timedelta(time_table["Minute"],unit="m")
    )
    return time_table

# ------ 修补time函数 ------
def repair_time_table(time_table:pd.DataFrame) -> pd.DataFrame:
    time_table = time_table.copy()

    time_table["TimeIndex"] = pd.to_numeric(
        time_table["TimeIndex"],
        errors="coerce"
    )
    time_table = time_table.dropna(subset=["TimeIndex","datetime"]) # subset为子集的意思，之看这两列，只要有NaN就整行删除
    time_table["TimeIndex"] = time_table["TimeIndex"].astype(int)

    start_index = int(time_table["TimeIndex"].min())
    end_index = int(time_table["TimeIndex"].max())

    full_time_index = pd.DataFrame(
        {
            "TimeIndex":range(start_index,end_index + 1)
        }
    )

    repaired_time_table = full_time_index.merge(
        time_table[["TimeIndex","datetime"]],
        on="TimeIndex",
        how="left",
    )
    first_row = time_table.sort_values("TimeIndex").iloc[0]
    first_time_index = int(first_row["TimeIndex"])
    first_datetime = pd.to_datetime(first_row["datetime"])

    repaired_time_table["datetime"] =(
        first_datetime
        + pd.to_timedelta(
            repaired_time_table["TimeIndex"] - first_time_index,
            unit="h",
        )
    )
# repaired_time_table["TimeIndex"] 是一个序列：[1, 2, 3, 4, 5, ...]

# first_time_index 是一个单个数字（比如 1）。

# Pandas 的规则：当“一列数据”与“单个数字”运算时，把那个数字复制成跟列一样长的副本，然后第 1 行减第 1 行，第 2 行减第 2 行……

# 结果得到新的序列：[0, 1, 2, 3, 4, ...]
    is_midnight = (
        (repaired_time_table["datetime"].dt.hour == 0)
        &(repaired_time_table["datetime"].dt.minute == 0)
    )

    energyplus_date = repaired_time_table["datetime"].where(
        ~is_midnight,
        repaired_time_table["datetime"] -pd.Timedelta(days=1)
    )
    repaired_time_table["Year"] = energyplus_date.dt.year
    repaired_time_table["Month"]= energyplus_date.dt.month
    repaired_time_table["Day"] = energyplus_date.dt.day
    repaired_time_table["Hour"] = repaired_time_table["datetime"].dt.hour
    repaired_time_table.loc[is_midnight,"Hour"] = 24
    repaired_time_table["Minute"] = repaired_time_table["datetime"].dt.minute

    return repaired_time_table[
        [
            "TimeIndex",
            "Year",
            "Month",
            "Day",
            "Hour",
            "Minute",
            "datetime",
        ]
    ]

# ------ 输出逐时能耗表 ------
def create_hourly_energy_table(
        report_data:pd.DataFrame,
        variable_name:str,
) -> pd.DataFrame:
    hourly_data = (
        report_data.loc[
            report_data["Name"] == variable_name,
            ["TimeIndex","datetime","Value"],
        ]
        .groupby(
            ["TimeIndex","datetime"],
            as_index=False
        )["Value"]
        .sum()
        .rename(columns={"Value":"Energy_J"}) # 针对Value列改名为Energy_j
    )
    hourly_data["Energy_KWh"] = (
        hourly_data["Energy_J"] / 3_600_000
    )

    return hourly_data

# ------ 计算单个变量的全年累计热量 ------
def calculate_annual_energy_kwh(
        report_data:pd.DataFrame,
        variable_name:str,
) -> float:
    hourly_data = create_hourly_energy_table(
        report_data=report_data,
        variable_name=variable_name,
    )

    if hourly_data.empty:
        raise ValueError(
            f"No report data found for variable: {variable_name}"
        )

    return float(hourly_data["Energy_KWh"].sum())
# ------ 计算各个能耗对总能耗的贡献 ------
def create_energy_contribution_table(
        report_data: pd.DataFrame,
) -> pd.DataFrame:
    contribution_records = []

    for component,variable_name in HEAT_CONTRIBUTION_VARIABLES.items():
        annual_energy_kwh = calculate_annual_energy_kwh(
            report_data=report_data,
            variable_name=variable_name,
        )

        contribution_records.append({
            "component": component,
            "variable_name": variable_name,
            "annual_energy_kwh": annual_energy_kwh,
        })

    contribution_table = pd.DataFrame(contribution_records)

    total_energy_kwh = contribution_table["annual_energy_kwh"].sum()

    if total_energy_kwh <= 0:
        raise ValueError(
            "Total heat contribution must be greater than zero."
        )

    contribution_table["contribution_percent"] = (
        contribution_table["annual_energy_kwh"]
        / total_energy_kwh
        *100
    )

    return contribution_table

# ------ 根据上面的逐时能耗表计算全年总能耗，最大逐时能耗以及峰值出现时间 ------
def calculate_energy_metrics(
        hourly_data: pd.DataFrame,
        interval_minutes : int
        ) -> dict:
    annual_energy_Kwh = hourly_data["Energy_KWh"].sum()

    peak_row_index = hourly_data["Energy_KWh"].idxmax()
    peak_row = hourly_data.loc[peak_row_index]

    peak_datetime_end = peak_row["datetime"]
    peak_datetime_start = (
        peak_datetime_end
        - pd.Timedelta(minutes=interval_minutes)
    )

    return {
        "annual_energy_kwh": annual_energy_Kwh,
        "peak_hourly_energy_kwh": peak_row["Energy_KWh"],
        "peak_datetime_start": peak_datetime_start,
        "peak_datetime_end": peak_datetime_end,
    }

# ------ 创建指标表 ------
def create_basic_metrics_table(
        cooling_metrics: dict,
        heating_metrics: dict,
) -> pd.DataFrame:
    metric_records = [
        {
            "metric": "cooling_demand",
            "annual_energy_kwh": float(
                cooling_metrics["annual_energy_kwh"]
            ),
            "peak_hourly_energy_kwh": float(
                cooling_metrics["peak_hourly_energy_kwh"]
            ),
            "peak_datetime_start": cooling_metrics["peak_datetime_start"],
            "peak_datetime_end": cooling_metrics["peak_datetime_end"],
        },
        {
            "metric": "heating_demand",
            "annual_energy_kwh": float(
                heating_metrics["annual_energy_kwh"]
            ),
            "peak_hourly_energy_kwh": float(
                heating_metrics["peak_hourly_energy_kwh"]
            ),
            "peak_datetime_start": heating_metrics["peak_datetime_start"],
            "peak_datetime_end": heating_metrics["peak_datetime_end"],
        },
    ]  # 按行加入Pandas Dataframe，一个集合里面是一行

    return pd.DataFrame(metric_records)

# ------ 输出以月份为单位的各项指标表 ------
def create_monthly_energy_table(
        cooling_hourly: pd.DataFrame,
        heating_hourly: pd.DataFrame,
 ) -> pd.DataFrame:
    cooling_data = cooling_hourly.copy()
    heating_data = heating_hourly.copy()

    cooling_data["month"] = cooling_data["datetime"].dt.month
    heating_data["month"] = heating_data["datetime"].dt.month

    monthly_cooling = (
        cooling_data
        .groupby(
            "month",
            as_index=False
        )["Energy_KWh"]
        .sum()
        .rename(
            columns = {
                "Energy_KWh" : "cooling_kwh"
            }
        )
    )

    monthly_heating = (
        heating_data
        .groupby(
            "month",
            as_index=False,
        )["Energy_KWh"]
        .sum()
        .rename(
            columns={
                "Energy_KWh" : "heating_kwh"
            }
        )
    )

    monthly_energy = monthly_cooling.merge(
        monthly_heating,
        on="month",
        how="outer"
    )

    monthly_energy = monthly_energy.sort_values(
        "month",
        ignore_index=True,
    ) 

    return monthly_energy

# ------ 峰值日筛选 ------
def extract_peak_day_profile(
        hourly_data:pd.DataFrame,
        peak_datetime:pd.Timestamp,
) -> pd.DataFrame:
    peak_date = peak_datetime.date() # 把日期加时刻的数据，只保留日期

    profile_data = hourly_data.copy()

    profile_data["interval_start"] = (
        profile_data["datetime"]
        - pd.Timedelta(hours=1)
    )
    peak_day_data = profile_data[
        profile_data["interval_start"].dt.date == peak_date
    ].copy()

    if len(peak_day_data) != 24:
        raise ValueError(
            f"Expected 24 hourly records for {peak_date}, "
            f"but found {len(peak_day_data)}."
        )

    peak_day_data = peak_day_data[
        [
           "interval_start",
            "datetime",
            "Energy_KWh", 
        ]
    ].copy()
    peak_day_data = peak_day_data.rename(
        columns={
            "Energy_KWh" : "energy_kwh"
        }
    )

    return peak_day_data

# ------ 表格增加峰值日的信息 ------
def create_peak_day_energy_table(
        cooling_peak_day:pd.DataFrame,
        heating_peak_day:pd.DataFrame,
) -> pd.DataFrame:
    cooling_data = cooling_peak_day.copy()

    cooling_data["load_type"] = "cooling"
    cooling_data["peak_date"] = (
        cooling_data["interval_start"].dt.date
    )
    cooling_data["hour"] = (
        cooling_data["interval_start"].dt.hour
    )

    heating_data = heating_peak_day.copy()
    heating_data["load_type"] = "heating"
    heating_data["peak_date"] = (
        heating_data["interval_start"].dt.date
    )
    heating_data["hour"] = (
        heating_data["interval_start"].dt.hour
    )

    peak_day_energy = pd.concat(
        [
            cooling_data,
            heating_data
        ],
        ignore_index=True
    )
    peak_day_energy = peak_day_energy[
        [
            "load_type",
            "peak_date",
            "hour",
            "interval_start",
            "datetime",
            "energy_kwh",
        ]
    ]

    return peak_day_energy
    
def build_overview_text(
    input_path: Path,
    sheet_names: list[str],
    dictionary_df: pd.DataFrame,
    time_df: pd.DataFrame,
    key_variables: pd.DataFrame,
    missing_sheets: list[str],
    missing_time_indices: list[int],
) -> str:
    lines = []

    lines.append("Building Energy Result Analyzer - Workbook Overview")
    lines.append("=" * 60)
    lines.append("")
    lines.append(f"Input file: {input_path}")
    lines.append(f"Total sheets: {len(sheet_names)}")
    lines.append("")
    lines.append("Sheets:")
    for sheet_name in sheet_names:
        lines.append(f"- {sheet_name}")

    lines.append("")
    lines.append("Required Sheet Check:")
    if missing_sheets:
        lines.append("Missing sheets:")
        for sheet_name in missing_sheets:
            lines.append(f"- {sheet_name}")
    else:
        lines.append("All required sheets were found.")

    lines.append("")
    lines.append("Main Tables:")
    lines.append(f"- ReportDataDictionary rows: {len(dictionary_df)}")
    lines.append(f"- Time rows: {len(time_df)}")
    lines.append(f"- Key variables found: {len(key_variables)}")

    lines.append("")
    lines.append("TimeIndex Check:")
    if missing_time_indices:
        preview = missing_time_indices[:20]
        lines.append(f"Missing TimeIndex values found: {preview}")
        if len(missing_time_indices) > 20:
            lines.append(f"... and {len(missing_time_indices) - 20} more")
    else:
        lines.append("No missing TimeIndex values found.")

    lines.append("")
    lines.append("Key Variables:")
    if key_variables.empty:
        lines.append("No key variables were found.")
    else:
        for _, row in key_variables.iterrows():
            lines.append(
                f"- {row['ReportDataDictionaryIndex']} | "
                f"{row['Name']} | "
                f"{row['KeyValue']} | "
                f"{row['Units']}"
            )

    lines.append("")

    return "\n".join(lines)

# ------ 主要运行阶段 ------
def run_analysis(input_path,output_dir):
    input_path = Path(input_path)
    output_dir = Path(output_dir)

    output_dir.mkdir(parents=True,exist_ok=True)

    workbook = load_workbook(
        input_path,
        read_only=True,
        data_only=True,
    )

    sheet_names = workbook.sheetnames
    missing_sheets = validate_sheets(sheet_names)

    dictionary_df = read_sheet_as_dataframe(workbook, "ReportDataDictionary")
    time_df = read_sheet_as_dataframe(workbook, "Time")

# ------ 获得时间戳及修复 ------

    time_table = prepare_time_table(time_df)
    missing_time_indices = find_missing_time_indices(time_df)
    print(f"Missing TimeIndex before repair: {missing_time_indices}")

    if missing_time_indices:
        time_table = repair_time_table(time_table)
        print(f"Time rows after repair: {len(time_table)}")
    else:
        print("Time table is complete. No repair needed.")

# ------ 查找关键参数index ------
    key_variables = find_key_variables(dictionary_df)
    key_variables_indices = get_key_variable_indices(key_variables)
    print(f"Key variable count:{len(key_variables_indices)}")

# ------ 将“指定sheet页”关键指标的数据抽出为dataframe ------
    selected_report_data_list = []
    for sheet_name in REPORT_DATA_SHEETS:
        selected_report_data = extract_selected_report_data_from_sheet(
            workbook=workbook,
            sheet_name=sheet_name,
            key_variable_indices=key_variables_indices
        )
        selected_report_data["SourceSheet"] = sheet_name
        print(f"Selected rows from {sheet_name}: {len(selected_report_data)}")

        selected_report_data_list.append(selected_report_data)

    selected_report_data =pd.concat(  # 将列表中的所有dataframe从上到下叠罗汉合并
        selected_report_data_list,
        ignore_index=True
    )

# ------ 将简单汇报表格加入值，名字，单位等信息 ------
    selected_report_data = attach_variable_info(
        report_data=selected_report_data,
        key_variables=key_variables,
    )
    selected_report_data = attach_time_info(
        report_data=selected_report_data,
        time_table=time_table,
    )

    missing_datetime_count = selected_report_data["datetime"].isna().sum()
    print(f"Rows without datetime:{missing_datetime_count}")

# ------ 输出逐时能耗表 ------
    cooling_hourly = create_hourly_energy_table(
        report_data=selected_report_data,
        variable_name=COOLING_NAME,
    )
    cooling_metrics = calculate_energy_metrics(
        hourly_data=cooling_hourly,
        interval_minutes=60,
    )

    heating_hourly = create_hourly_energy_table(
        report_data=selected_report_data,
        variable_name=HEATING_NAME
    )
    heating_metrics = calculate_energy_metrics(
        hourly_data=heating_hourly,
        interval_minutes=60,
    )

# ------ 输出冷热峰值日 ------
    cooling_peak_day = extract_peak_day_profile(
        hourly_data=cooling_hourly,
        peak_datetime=cooling_metrics["peak_datetime_start"],
    )
    heating_peak_day = extract_peak_day_profile(
        hourly_data=heating_hourly,
        peak_datetime=heating_metrics["peak_datetime_start"],
    )

    print(
        f"Cooling peak day rows: {len(cooling_peak_day)}"
    )
    print(
        f"Heating peak day rows: {len(heating_peak_day)}"
    )
# ------ 丰富峰值日指标 ------
    peak_day_energy = create_peak_day_energy_table(
        cooling_peak_day=cooling_peak_day,
        heating_peak_day=heating_peak_day,
    )

    peak_day_energy_path = output_dir / "peak_day_energy.csv"
    peak_day_energy.to_csv(
        peak_day_energy_path,
        index=False,
        encoding="utf-8-sig",
    )
    print("Peak day energy:")
    print(peak_day_energy.to_string(index=False))
    print(f"Peak day energy rows: {len(peak_day_energy)}")
    print(f"Peak day energy saved to: {peak_day_energy_path}")

# ------ 输出月度能耗表 ------
    monthly_energy = create_monthly_energy_table(
        cooling_hourly = cooling_hourly,
        heating_hourly = heating_hourly,
    )

    print(cooling_hourly.head())
    print(f"Cooling hourly rows:{len(cooling_hourly)}")
    print("Cooling metrics:")
    print(cooling_metrics)
    print("Heating metrics:")
    print(heating_metrics)

# ------ 输出冷热各项指标表 ------
    basic_metrics = create_basic_metrics_table(
        cooling_metrics=cooling_metrics,
        heating_metrics=heating_metrics,
    )

    basic_metrics_path = output_dir / "basic_metrics.csv"
    basic_metrics.to_csv(
        basic_metrics_path,
        index=False,
        encoding="utf-8-sig",
    )

    print("Basic energy metrics:")
    print(basic_metrics.to_string(index=False))
    print(f"Basic metrics saved to: {basic_metrics_path}")

# ------ 输出月度能耗表 ------
    monthly_energy_path = output_dir / "monthly_energy.csv"

    monthly_energy.to_csv(
        monthly_energy_path,
        index=False,
        encoding = "utf-8-sig",
    )
  
    print("Monthly energy:")
    print(monthly_energy.to_string(index=False))
    print(f"Monthly energy saved to: {monthly_energy_path}")
# ------ 输出内部得热与太阳得热组成表 ------
    energy_contribution = create_energy_contribution_table(
        report_data=selected_report_data,
    )

    energy_contribution_path = output_dir / "energy_contribution.csv"
    energy_contribution.to_csv(
        energy_contribution_path,
        index=False,
        encoding="utf-8-sig"
    )

    print("Energy contribution:")
    print(energy_contribution.to_string(index=False))
    print(
        "Contribution percent total:"
        f"{energy_contribution["contribution_percent"].sum():.2f}%"
    )
    print(
        f"Energy contribution saved to: {energy_contribution_path}"
    )

# ------ 把已经筛选并合并好变量名称的数据，保存成 CSV 文件 ------
    selected_report_data_path = output_dir / "selected_report_data.csv"
    selected_report_data.to_csv( selected_report_data_path,
    index=False,
    encoding="utf-8-sig",)
    print(f"Selected report data saved to: {selected_report_data_path}")
    print(f"Total selected rows: {len(selected_report_data)}")
# ------ 查找缺少的时间戳 ------

    missing_time_indices = find_missing_time_indices(time_df)

# ------ 撰写预览信息 ------
    overview_text = build_overview_text(
        input_path=input_path,
        sheet_names=sheet_names,
        dictionary_df=dictionary_df,
        time_df=time_df,
        key_variables=key_variables,
        missing_sheets=missing_sheets,
        missing_time_indices=missing_time_indices,
    )

    overview_path = output_dir / "workbook_overview.txt"
    key_variables_path = output_dir / "key_variables.csv"

    overview_path.write_text(overview_text, encoding="utf-8")
    key_variables.to_csv(key_variables_path, index=False, encoding="utf-8-sig")

    print("Workbook inspection finished.")
    print(f"Overview saved to: {overview_path}")
    print(f"Key variables saved to: {key_variables_path}")
    print(f"Total sheets: {len(sheet_names)}")
    print(f"Dictionary rows: {len(dictionary_df)}")
    print(f"Time rows: {len(time_df)}")
    print(f"Key variables found: {len(key_variables)}")

def main() -> None:
    args = parse_arg()

    input_path = resolve_path(args.input)
    output_dir = resolve_path(args.output)

    if not input_path.exists():
        print(f"Input file not found: {input_path}")
        print("Please put run_sql_export.xlsx into the data folder first.")
        sys.exit(1)

    run_analysis(
        input_path,
        output_dir
    )

if __name__ == "__main__":
    main()